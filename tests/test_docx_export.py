import os
import sys
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from docx import Document

from src.export.docx_export import render_letter_docx

DRAFT = """सारण समाहरणालय, छपरा

विषय: शिक्षकों की कमी के संबंध में।

महोदय,

निवेदन है कि शिक्षकों के पद रिक्त हैं।

भवदीय,
[नाम]"""


def test_render_letter_docx_creates_file(tmp_path: Path):
    output = render_letter_docx(
        DRAFT,
        tmp_path / "letter.docx",
        font_name="Noto Sans Devanagari",
        font_size_pt=12,
        margin_cm=2.5,
        line_spacing=1.5,
    )
    assert output.exists()
    assert output.suffix == ".docx"


def test_render_letter_docx_preserves_all_text(tmp_path: Path):
    output = render_letter_docx(
        DRAFT,
        tmp_path / "letter.docx",
        font_name="Noto Sans Devanagari",
        font_size_pt=12,
        margin_cm=2.5,
        line_spacing=1.5,
    )
    doc = Document(str(output))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "शिक्षकों की कमी के संबंध में" in full_text
    assert "[नाम]" in full_text


def test_render_letter_docx_sets_font_and_margins(tmp_path: Path):
    output = render_letter_docx(
        DRAFT,
        tmp_path / "letter.docx",
        font_name="Noto Sans Devanagari",
        font_size_pt=14,
        margin_cm=3.0,
        line_spacing=1.5,
    )
    doc = Document(str(output))
    section = doc.sections[0]
    assert round(section.left_margin.cm, 1) == 3.0

    first_run = doc.paragraphs[0].runs[0]
    assert first_run.font.name == "Noto Sans Devanagari"
    assert first_run.font.size.pt == 14


def test_render_letter_docx_creates_parent_dirs(tmp_path: Path):
    nested = tmp_path / "a" / "b" / "letter.docx"
    render_letter_docx(
        DRAFT, nested,
        font_name="Noto Sans Devanagari", font_size_pt=12,
        margin_cm=2.5, line_spacing=1.5,
    )
    assert nested.exists()
